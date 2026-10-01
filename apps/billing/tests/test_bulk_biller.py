"""
Tests para `services.bulk_biller` y la vista de emisión masiva.

No llaman a la API real: el parseo de la planilla (SDK), `biller.draft()`/
`bill_draft()` y el render de PDF se mockean siempre, y la tarea en segundo
plano no se encola de verdad. Lo que se prueba es lo que hace Slim (elegir la
estrategia, mostrar el error del parseo en el formulario, guardar los
documentos, emitirlos, armar el resultado y el correo), no el formato de la
planilla ni los cálculos, que viven en la biblioteca.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from unittest import mock

import pytest
import yaml
from django.contrib.auth.models import User
from django.core import mail
from django.core.cache import cache
from django.test import Client
from django.urls import reverse
from kombu.exceptions import OperationalError
from libredte_lib_sdk.exceptions import LibreDteApiError
from openpyxl import load_workbook

from apps.billing import tasks
from apps.billing.models import Borrador, TipoDte
from apps.billing.services import bulk_biller
from apps.libredte import tenancy
from apps.libredte.models import Comuna, Contribuyente, Pais
from apps.libredte.services.exceptions import ServiceError

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def cache_local(settings: Any) -> Iterator[None]:
    """
    Caché en memoria y vacía, en vez del Redis de `CACHES`.

    La usan los bloqueos de la emisión masiva: sin esto los tests
    necesitarían un Redis levantado, y un bloqueo de un test
    (tomado por la vista y nunca liberado, porque la tarea está
    mockeada) rechazaría la misma planilla en el test siguiente.
    """
    settings.CACHES = {
        'default': {
            'BACKEND': 'django.core.cache.backends.locmem.LocMemCache',
        },
    }
    cache.clear()
    yield
    cache.clear()


@pytest.fixture(autouse=True)
def carpeta_pdf(settings: Any, tmp_path: Path) -> Path:
    """Los ZIP de PDF van a una carpeta temporal del test."""
    carpeta = tmp_path / 'emision_masiva_pdf'
    settings.EMISION_MASIVA_PDF_DIR = carpeta
    return carpeta


@pytest.fixture(autouse=True)
def chile() -> Pais:
    return Pais.objects.create(codigo=Pais.CHILE, glosa='Chile')


@pytest.fixture
def comuna() -> Comuna:
    return Comuna.objects.create(codigo=13101, glosa='Santiago')


@pytest.fixture
def usuario() -> User:
    return User.objects.create_user(
        'demo', email='demo@example.com', password='demo12345'
    )


@pytest.fixture
def contribuyente(comuna: Comuna, usuario: User) -> Contribuyente:
    return Contribuyente.objects.create(
        usuario=usuario,
        rut=76192083,
        dv='9',
        razon_social='SASCO SpA',
        giro='Servicios',
        direccion='Av. Siempre Viva 123',
        comuna=comuna,
        autorizacion_dte_resolucion_fecha='2014-08-22',
        autorizacion_dte_resolucion_numero=0,
    )


@pytest.fixture(autouse=True)
def tipo_factura() -> TipoDte:
    return TipoDte.objects.create(codigo=33, glosa='Factura Electrónica')


def _documento(
    folio: int = 1,
    *,
    tipo: int = 33,
    nombre_pdf: str | None = None,
) -> dict[str, Any]:
    """Un documento tal como lo entrega la biblioteca (`document_parsed`)."""
    datos: dict[str, Any] = {
        'Encabezado': {
            'IdDoc': {'TipoDTE': tipo, 'Folio': folio},
            'Receptor': {
                'RUTRecep': '55666777-2',
                'RznSocRecep': 'Cliente de prueba',
            },
        },
        'Detalle': [
            {'NmbItem': 'Servicio web', 'QtyItem': 1, 'PrcItem': 6000}
        ],
    }
    if nombre_pdf is not None:
        datos['LibreDTE'] = {'pdf': {'nombre': nombre_pdf}}
    return datos


def _backend(
    *documentos: dict[str, Any],
    error: Exception | None = None,
) -> mock.MagicMock:
    """El backend del SDK, con `batch_processor.parse` ya resuelto."""
    backend = mock.MagicMock()
    parse = backend.billing.document.batch_processor.parse
    if error is not None:
        parse.side_effect = error
    else:
        parse.return_value = mock.Mock(
            document_bags=[
                mock.Mock(document_parsed=documento)
                for documento in documentos
            ],
        )
    return backend


def _parsear(
    contribuyente: Contribuyente,
    backend: mock.MagicMock,
    extension: str = 'xlsx',
) -> list[dict[str, Any]]:
    """`bulk_biller.parse()` con el backend dado."""
    with mock.patch(
        'apps.billing.services.bulk_biller.get_backend'
    ) as get_backend:
        get_backend.return_value.__enter__.return_value = backend
        return bulk_biller.parse(contribuyente, b'contenido', extension)


@pytest.mark.parametrize(
    ('extension', 'estrategia'),
    [('csv', 'spreadsheet.csv'), ('xlsx', 'spreadsheet.xlsx')],
)
def test_el_formato_elige_la_estrategia_de_la_biblioteca(
    contribuyente: Contribuyente,
    extension: str,
    estrategia: str,
) -> None:
    backend = _backend(_documento(1), _documento(2))

    documentos = _parsear(contribuyente, backend, extension)

    backend.billing.document.batch_processor.parse.assert_called_once_with(
        b'contenido', strategy=estrategia
    )
    assert [d['Encabezado']['IdDoc']['Folio'] for d in documentos] == [1, 2]


def test_un_formato_no_soportado_no_llega_a_la_biblioteca(
    contribuyente: Contribuyente,
) -> None:
    backend = _backend(_documento())

    with pytest.raises(ServiceError, match='CSV o XLSX'):
        _parsear(contribuyente, backend, 'pdf')

    backend.billing.document.batch_processor.parse.assert_not_called()


def test_el_error_de_la_biblioteca_sale_con_su_fila(
    contribuyente: Contribuyente,
) -> None:
    error = LibreDteApiError(
        operation_id='billing.document.batch_processor::parse',
        status_code=500,
        title='BatchProcessorException',
        detail='Fila 3: Falta RUT del receptor.',
    )

    with pytest.raises(ServiceError) as excinfo:
        _parsear(contribuyente, _backend(error=error))

    # Solo el detalle, sin el nombre de la clase de la excepción.
    assert str(excinfo.value) == 'Fila 3: Falta RUT del receptor.'


def test_los_documentos_se_guardan_en_disco_con_su_huella(
    contribuyente: Contribuyente,
    carpeta_pdf: Path,
) -> None:
    documentos = [_documento(1), _documento(2)]

    ruta = bulk_biller.save_documentos(contribuyente, documentos)

    assert ruta.parent == carpeta_pdf / '76192083' / 'lotes'
    assert bulk_biller.load_documentos(ruta) == documentos
    # La misma planilla cae en el mismo archivo; otra, en otro.
    assert bulk_biller.save_documentos(contribuyente, documentos) == ruta
    assert bulk_biller.save_documentos(contribuyente, [_documento(3)]) != ruta


def _autenticar(client: Client, contribuyente: Contribuyente) -> None:
    """Login + deja `contribuyente` activo en la sesión del cliente."""
    client.force_login(contribuyente.usuario)
    session = client.session
    session[tenancy._SESSION_KEY] = contribuyente.pk
    session.save()


_PARSE = 'apps.billing.services.bulk_biller.parse'
_DELAY = 'apps.billing.tasks.bulk_biller.process_bulk_billing.delay'


def _subir(
    client: Client,
    nombre: str = 'documentos.xlsx',
    **opciones: str,
) -> Any:
    """POST de una planilla a `/billing/emitir/masivo`."""
    planilla = io.BytesIO(b'contenido de la planilla')
    planilla.name = nombre
    return client.post(
        reverse('billing:emitir_masivo'),
        {'archivo': planilla, 'modo': 'borrador', 'pdf': '', **opciones},
    )


def test_vista_encola_los_documentos_ya_parseados(
    client: Client,
    contribuyente: Contribuyente,
    carpeta_pdf: Path,
) -> None:
    """
    La planilla se parsea en el request. Lo que se encola es la ruta de un
    archivo con los documentos, no la planilla.
    """
    _autenticar(client, contribuyente)
    documentos = [_documento(1), _documento(2)]

    with mock.patch(_PARSE, return_value=documentos) as parse:
        with mock.patch(_DELAY) as delay:
            response = _subir(client, modo='real', pdf='1')

    parse.assert_called_once_with(
        contribuyente, b'contenido de la planilla', 'xlsx'
    )
    assert response.status_code == 302
    assert response['Location'] == reverse('billing:emitir_masivo')
    delay.assert_called_once()
    args, kwargs = delay.call_args
    assert args[:2] == (contribuyente.pk, contribuyente.usuario.pk)
    assert kwargs == {'dte_real': True, 'pdf': True}
    assert bulk_biller.load_documentos(Path(args[2])) == documentos
    avisos = [str(m) for m in response.wsgi_request._messages]
    assert any('demo@example.com' in aviso for aviso in avisos)


def test_vista_muestra_el_error_del_parseo_en_el_formulario(
    client: Client,
    contribuyente: Contribuyente,
    carpeta_pdf: Path,
) -> None:
    """Se informa el primer error y no se guarda ni se encola nada."""
    _autenticar(client, contribuyente)

    with mock.patch(
        _PARSE, side_effect=ServiceError('Fila 3: Falta RUT del receptor.')
    ):
        with mock.patch(_DELAY) as delay:
            response = _subir(client)

    assert response.status_code == 200
    assert response.context['form'].errors['archivo'] == [
        'Fila 3: Falta RUT del receptor.'
    ]
    delay.assert_not_called()
    assert not carpeta_pdf.exists()


def test_vista_rechaza_una_planilla_sin_documentos(
    client: Client,
    contribuyente: Contribuyente,
    carpeta_pdf: Path,
) -> None:
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, return_value=[]), mock.patch(_DELAY) as delay:
        response = _subir(client)

    assert response.status_code == 200
    assert response.context['form'].errors['archivo'] == [
        'La planilla no tiene documentos.'
    ]
    delay.assert_not_called()
    assert not carpeta_pdf.exists()


def test_solo_ver_los_documentos_descarga_el_yaml_sin_emitir_nada(
    client: Client,
    contribuyente: Contribuyente,
    carpeta_pdf: Path,
) -> None:
    """Sirve para ver lo que se generará: no guarda ni encola nada."""
    _autenticar(client, contribuyente)
    documentos = [_documento(1), _documento(2)]

    with mock.patch(_PARSE, return_value=documentos):
        with mock.patch(_DELAY) as delay:
            response = _subir(client, modo='ver')

    assert response.status_code == 200
    assert response['Content-Type'].startswith('application/yaml')
    assert 'documentos_documentos.yaml' in response['Content-Disposition']
    assert yaml.safe_load(response.content) == documentos
    delay.assert_not_called()
    assert not carpeta_pdf.exists()
    assert mail.outbox == []


def test_solo_ver_los_documentos_no_necesita_correo(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    """No se emite ni se envía nada: el correo del usuario no hace falta."""
    contribuyente.usuario.email = ''
    contribuyente.usuario.save()
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, return_value=[_documento()]):
        response = _subir(client, modo='ver')

    assert response.status_code == 200
    assert response['Content-Type'].startswith('application/yaml')


def test_solo_ver_los_documentos_tambien_muestra_el_error_del_parseo(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, side_effect=ServiceError('Fila 3: sin RUT.')):
        response = _subir(client, modo='ver')

    assert response.status_code == 200
    assert response.context['form'].errors['archivo'] == ['Fila 3: sin RUT.']


def test_el_modo_por_defecto_es_borradores_y_hay_tres(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    _autenticar(client, contribuyente)

    response = client.get(reverse('billing:emitir_masivo'))

    campo = response.context['form']['modo']
    assert campo.initial == 'borrador'
    assert [valor for valor, _ in campo.field.choices] == [
        'ver',
        'borrador',
        'real',
    ]
    assert 'Modo de ejecución' in response.content.decode()


def test_vista_rechaza_un_modo_que_no_existe(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE) as parse, mock.patch(_DELAY) as delay:
        response = _subir(client, modo='cualquiera')

    assert response.status_code == 200
    assert response.context['form'].errors['modo']
    parse.assert_not_called()
    delay.assert_not_called()


def test_el_modo_borradores_no_emite_el_dte_real(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, return_value=[_documento()]):
        with mock.patch(_DELAY) as delay:
            _subir(client, modo='borrador')

    assert delay.call_args.kwargs['dte_real'] is False


def test_vista_rechaza_una_extension_que_no_es_csv_ni_xlsx(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    _autenticar(client, contribuyente)

    with mock.patch(_DELAY) as delay:
        response = _subir(client, nombre='documentos.pdf')

    assert response.status_code == 200
    assert 'CSV o XLSX' in response.context['form'].errors['archivo'][0]
    delay.assert_not_called()


def test_vista_no_encola_si_el_usuario_no_tiene_correo(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    """Sin correo, nadie se enteraría de cómo terminó la emisión."""
    contribuyente.usuario.email = ''
    contribuyente.usuario.save()
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE) as parse, mock.patch(_DELAY) as delay:
        response = _subir(client)

    assert response.status_code == 200
    parse.assert_not_called()
    delay.assert_not_called()
    assert response.context['form'].non_field_errors()


def test_vista_avisa_si_no_pudo_encolar(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    """Con el broker caído, no se da la planilla por encolada."""
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, return_value=[_documento()]):
        with mock.patch(_DELAY, side_effect=OperationalError('sin redis')):
            response = _subir(client)

    assert response.status_code == 302
    avisos = [str(m) for m in response.wsgi_request._messages]
    assert any('No fue posible programar' in aviso for aviso in avisos)


def test_la_pagina_explica_el_formato_y_enlaza_los_ejemplos(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    """
    La tabla de códigos y las reglas de la planilla son referencia para
    llenarla, así que están desde el `GET`, no recién junto con el
    resultado.
    """
    _autenticar(client, contribuyente)

    response = client.get(reverse('billing:emitir_masivo'))

    html = response.content.decode()
    for glosa in bulk_biller.GLOSAS_CODIGO.values():
        assert glosa in html
    assert 'las demás vacías' in html
    assert 'Se recomienda no más de 100 documentos' in html
    for extension in ('csv', 'xlsx'):
        enlace = reverse('billing:emitir_masivo_plantilla', args=[extension])
        assert enlace in html


@pytest.mark.parametrize('extension', ['csv', 'xlsx'])
def test_descarga_de_las_planillas_de_ejemplo(
    client: Client,
    contribuyente: Contribuyente,
    extension: str,
) -> None:
    _autenticar(client, contribuyente)

    response = client.get(
        reverse('billing:emitir_masivo_plantilla', args=[extension])
    )

    assert response.status_code == 200
    assert f'emision_masiva.{extension}' in response['Content-Disposition']


def test_no_hay_ejemplo_con_otra_extension(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    _autenticar(client, contribuyente)

    response = client.get(
        reverse('billing:emitir_masivo_plantilla', args=['pdf'])
    )

    assert response.status_code == 404


def test_la_misma_planilla_no_se_encola_dos_veces(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    """Un doble clic emitiría cada documento dos veces."""
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, return_value=[_documento()]):
        with mock.patch(_DELAY) as delay:
            _subir(client)
            response = _subir(client)

    delay.assert_called_once()
    avisos = [str(m) for m in response.wsgi_request._messages]
    assert any('ya se está procesando' in aviso for aviso in avisos)


def test_si_no_pudo_encolar_la_planilla_se_puede_reintentar(
    client: Client,
    contribuyente: Contribuyente,
) -> None:
    """Un encolado fallido no deja la planilla bloqueada ni su archivo."""
    _autenticar(client, contribuyente)

    with mock.patch(_PARSE, return_value=[_documento()]):
        with mock.patch(
            _DELAY, side_effect=[OperationalError('sin redis'), None]
        ) as delay:
            _subir(client)
            response = _subir(client)

    assert delay.call_count == 2
    avisos = [str(m) for m in response.wsgi_request._messages]
    assert any('está siendo procesada' in aviso for aviso in avisos)


def _emitir(
    contribuyente: Contribuyente,
    usuario: User,
    *documentos: dict[str, Any],
    dte_real: bool = False,
) -> tuple[list[bulk_biller.Resultado], mock.MagicMock, mock.MagicMock]:
    """Corre `bulk_biller.process()` con el `biller` mockeado."""
    with (
        mock.patch('apps.billing.services.bulk_biller.biller.draft') as draft,
        mock.patch(
            'apps.billing.services.bulk_biller.biller.bill_draft'
        ) as bill,
    ):
        draft.return_value = mock.Mock(pk=7, codigo='33-0000007')
        bill.return_value = mock.Mock(pk=9, folio=42)
        emision = bulk_biller.process(
            contribuyente,
            list(documentos),
            usuario,
            dte_real=dte_real,
            pdf=False,
        )
    return emision.resultados, draft, bill


def test_un_documento_queda_como_borrador(
    contribuyente: Contribuyente,
    usuario: User,
    tipo_factura: TipoDte,
) -> None:
    documento = _documento()

    resultados, draft, bill = _emitir(contribuyente, usuario, documento)

    assert [r.ok for r in resultados] == [True]
    draft.assert_called_once_with(
        contribuyente, tipo_factura, documento, usuario
    )
    bill.assert_not_called()
    assert resultados[0].glosa == 'Borrador creado 33-0000007'


def test_dte_real_confirma_el_borrador(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    resultados, _, bill = _emitir(
        contribuyente, usuario, _documento(), dte_real=True
    )

    assert [r.ok for r in resultados] == [True]
    assert bill.call_count == 1
    assert '42' in resultados[0].glosa


def test_el_resultado_trae_el_numero_tipo_folio_y_receptor_del_documento(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    resultados, _, _ = _emitir(
        contribuyente, usuario, _documento(5), _documento(6)
    )

    assert [
        (r.numero, r.tipo_dte, r.folio_planilla, r.rut_receptor)
        for r in resultados
    ] == [(1, '33', '5', '55666777-2'), (2, '33', '6', '55666777-2')]


def test_error_de_la_api_no_frena_los_documentos_siguientes(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    with (
        mock.patch('apps.billing.services.bulk_biller.biller.draft') as draft,
        mock.patch('apps.billing.services.bulk_biller.biller.bill_draft'),
    ):
        draft.side_effect = [ServiceError('falta el giro'), mock.Mock()]
        resultados = bulk_biller.process(
            contribuyente,
            [_documento(1), _documento(2)],
            usuario,
            dte_real=False,
            pdf=False,
        ).resultados

    assert [r.codigo for r in resultados] == [
        bulk_biller.CODIGO_BORRADOR,
        None,
    ]
    assert resultados[0].glosa == 'falta el giro'


def test_un_tipo_de_documento_que_no_existe_es_error_de_ese_documento(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    resultados, draft, _ = _emitir(
        contribuyente, usuario, _documento(1, tipo=999), _documento(2)
    )

    assert [r.codigo for r in resultados] == [
        bulk_biller.CODIGO_BORRADOR,
        None,
    ]
    assert draft.call_count == 1


def _procesar(
    contribuyente: Contribuyente,
    usuario: User,
    *documentos: dict[str, Any],
    pdf: bool = False,
) -> None:
    """Corre `bulk_biller.process()` con la API mockeada."""
    borrador = mock.Mock(codigo='33-0000007')
    with (
        mock.patch(
            'apps.billing.services.bulk_biller.biller.draft',
            return_value=mock.Mock(pk=7, codigo='33-0000007'),
        ),
        mock.patch(
            'apps.billing.services.bulk_biller.Borrador.objects.get',
            return_value=borrador,
        ),
        mock.patch(
            'apps.billing.services.bulk_biller.document_renderer.render_pdf',
            return_value=mock.Mock(content_bytes=b'%PDF-1.4 falso'),
        ),
    ):
        bulk_biller.process(
            contribuyente,
            list(documentos),
            usuario,
            dte_real=False,
            pdf=pdf,
        )


def _adjunto(nombre_termina_en: str) -> bytes:
    """El contenido del adjunto del único correo enviado."""
    (mensaje,) = mail.outbox
    for nombre, contenido, _ in mensaje.attachments:
        if nombre.endswith(nombre_termina_en):
            assert isinstance(contenido, bytes)
            return contenido
    raise AssertionError(f'Sin adjunto {nombre_termina_en}')


def _token_del_correo() -> str:
    """El token del enlace de descarga de PDF del único correo enviado."""
    (mensaje,) = mail.outbox
    encontrado = re.search(r'/emitir/masivo/pdf/([^/\s]+)/', str(mensaje.body))
    assert encontrado is not None, 'El correo no trae el enlace de los PDF'
    return encontrado.group(1)


def _zip_del_correo() -> bytes:
    """El ZIP de PDF al que apunta el enlace del único correo enviado."""
    encontrado = bulk_biller.find_pdf_zip(_token_del_correo())
    assert encontrado is not None, 'El enlace de los PDF no sirve'
    ruta, _ = encontrado
    return ruta.read_bytes()


def test_procesar_avisa_por_correo_al_usuario(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """Solo al usuario que subió la planilla, nunca al receptor."""
    _procesar(contribuyente, usuario, _documento())

    (mensaje,) = mail.outbox
    assert mensaje.to == ['demo@example.com']
    assert 'Resultado emisión masiva de DTE' in mensaje.subject
    assert 'Documentos: 1 (0 con error)' in mensaje.body


def test_el_resultado_tiene_una_fila_por_documento(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """Es el archivo que va adjunto al correo."""
    _procesar(contribuyente, usuario, _documento(1), _documento(2))

    hoja = load_workbook(io.BytesIO(_adjunto('.xlsx'))).active
    assert hoja is not None
    filas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
    assert filas[0] == [
        'N°',
        'Tipo DTE',
        'Folio',
        'RUT Receptor',
        'Resultado código',
        'Resultado glosa',
    ]
    # Sin código: el documento se generó.
    assert filas[1] == [
        '1',
        '33',
        '1',
        '55666777-2',
        None,
        'Borrador creado 33-0000007',
    ]
    assert filas[2][:3] == ['2', '33', '2']
    assert len(filas) == 3


def test_los_pdf_van_en_un_zip_que_se_descarga_desde_el_correo(
    contribuyente: Contribuyente,
    usuario: User,
    carpeta_pdf: Path,
) -> None:
    _procesar(contribuyente, usuario, _documento(), pdf=True)

    # Una carpeta por contribuyente, con su RUT sin dígito verificador.
    (carpeta_contribuyente,) = carpeta_pdf.iterdir()
    assert carpeta_contribuyente.name == '76192083'

    with zipfile.ZipFile(io.BytesIO(_zip_del_correo())) as zip_pdf:
        assert zip_pdf.namelist() == ['LibreDTE_76192083_33-0000007.pdf']


def test_el_nombre_del_pdf_sale_del_documento(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """
    Admite `{rut}`, `{dv}`, `{dte}` y `{folio}`. Lo que no sea letra,
    número, punto o guion se reemplaza: una barra no puede crear
    carpetas dentro del ZIP.
    """
    _procesar(
        contribuyente,
        usuario,
        _documento(nombre_pdf='factura/{rut}-{dv}_{dte}_{folio}'),
        pdf=True,
    )

    with zipfile.ZipFile(io.BytesIO(_zip_del_correo())) as zip_pdf:
        assert zip_pdf.namelist() == ['factura_76192083-9_33_33-0000007.pdf']


def test_si_algo_falla_igual_se_avisa_por_correo(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """
    El usuario no tiene otra forma de enterarse de qué pasó con su
    planilla. Después de avisar, el error se vuelve a levantar para que
    la tarea quede registrada como fallida.
    """
    with (
        mock.patch(
            'apps.billing.services.bulk_biller._bill_documento',
            side_effect=RuntimeError('se cayó la base'),
        ),
        pytest.raises(RuntimeError),
    ):
        bulk_biller.process(
            contribuyente,
            [_documento()],
            usuario,
            dte_real=False,
            pdf=False,
        )

    (mensaje,) = mail.outbox
    assert 'se cayó la base' in mensaje.body


def test_la_tarea_devuelve_un_reporte_serializable(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """
    El reporte queda guardado como resultado de la tarea
    (`django_celery_results`), que lo serializa a JSON: un campo que no
    lo fuera haría fallar la tarea recién al guardar su resultado.
    """
    ruta = bulk_biller.save_documentos(contribuyente, [_documento()])

    with mock.patch(
        'apps.billing.services.bulk_biller.biller.draft',
        return_value=mock.Mock(pk=7, codigo='33-0000007'),
    ):
        reporte = tasks.process_bulk_billing(
            contribuyente.pk,
            usuario.pk,
            str(ruta),
            dte_real=False,
            pdf=False,
        )

    assert json.loads(json.dumps(reporte)) == reporte
    assert reporte['status'] is True
    assert str(reporte['envio']) in str(mail.outbox[0].subject)
    assert reporte['result'] == [
        {
            'numero': 1,
            'tipo_dte': '33',
            'folio_planilla': '1',
            'rut_receptor': '55666777-2',
            'codigo': None,
            'glosa': 'Borrador creado 33-0000007',
            'borrador_pk': 7,
            'emitido_pk': None,
        },
    ]


def test_la_tarea_borra_el_archivo_y_libera_el_bloqueo_aunque_falle(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """Terminada la tarea, la misma planilla se puede volver a subir."""
    ruta = bulk_biller.save_documentos(contribuyente, [_documento()])
    bulk_biller.acquire_lock(contribuyente.pk, ruta.stem)

    with (
        mock.patch(
            'apps.billing.services.bulk_biller.process',
            side_effect=RuntimeError('se cayó la base'),
        ),
        pytest.raises(RuntimeError),
    ):
        tasks.process_bulk_billing(
            contribuyente.pk,
            usuario.pk,
            str(ruta),
            dte_real=False,
            pdf=False,
        )

    assert not ruta.exists()
    # Si el bloqueo siguiera tomado, esto levantaría
    # `BulkBillingInProgressError`.
    bulk_biller.acquire_lock(contribuyente.pk, ruta.stem)


def test_si_falla_a_mitad_de_camino_llega_lo_ya_procesado(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """
    Sin el resultado de lo que alcanzó a emitirse, el usuario no sabría
    qué quitar de la planilla antes de volver a subirla.
    """
    primero = bulk_biller.Resultado(
        numero=1,
        tipo_dte='33',
        folio_planilla='1',
        rut_receptor='55666777-2',
        codigo=None,
        glosa='Borrador creado 33-0000007',
        borrador_pk=7,
    )
    with (
        mock.patch(
            'apps.billing.services.bulk_biller._bill_documento',
            side_effect=[primero, RuntimeError('se cayó la base')],
        ),
        pytest.raises(RuntimeError),
    ):
        bulk_biller.process(
            contribuyente,
            [_documento(1), _documento(2)],
            usuario,
            dte_real=False,
            pdf=False,
        )

    (mensaje,) = mail.outbox
    assert 'se cayó la base' in mensaje.body
    assert 'Documentos procesados: 1 (0 con error)' in mensaje.body
    hoja = load_workbook(io.BytesIO(_adjunto('.xlsx'))).active
    assert hoja is not None
    filas = [list(fila) for fila in hoja.iter_rows(values_only=True)]
    assert len(filas) == 2
    assert filas[1][-1] == 'Borrador creado 33-0000007'


def test_documento_borrado_antes_de_su_pdf_cuenta_como_fallido(
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    """Un PDF que falta no puede cortar el correo de toda la emisión."""
    with (
        mock.patch(
            'apps.billing.services.bulk_biller.biller.draft',
            return_value=mock.Mock(pk=7, codigo='33-0000007'),
        ),
        mock.patch(
            'apps.billing.services.bulk_biller.Borrador.objects.get',
            side_effect=Borrador.DoesNotExist,
        ),
    ):
        bulk_biller.process(
            contribuyente,
            [_documento()],
            usuario,
            dte_real=False,
            pdf=True,
        )

    (mensaje,) = mail.outbox
    assert 'no se pudieron generar 1' in mensaje.body
    with zipfile.ZipFile(io.BytesIO(_zip_del_correo())) as zip_pdf:
        assert zip_pdf.namelist() == []


def test_el_zip_no_va_adjunto_y_el_enlace_es_absoluto(
    contribuyente: Contribuyente,
    usuario: User,
    settings: Any,
) -> None:
    """
    El correo no aguantaría el ZIP de una planilla completa. El enlace
    se arma en el worker, sin request, así que lleva el dominio de
    `HOSTNAME`.
    """
    settings.URL_SCHEMA = 'https'
    settings.HOSTNAME = ['slim.example.com']

    _procesar(contribuyente, usuario, _documento(), pdf=True)

    (mensaje,) = mail.outbox
    ((nombre, _, _),) = mensaje.attachments
    assert nombre.endswith('.xlsx')
    assert 'https://slim.example.com/billing/emitir/masivo/pdf/' in (
        mensaje.body
    )
    assert 'Descargue los PDF en el siguiente enlace' in mensaje.body


def test_el_enlace_descarga_el_zip_sin_login(
    client: Client,
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    _procesar(contribuyente, usuario, _documento(), pdf=True)

    response = client.get(
        reverse(
            'billing:emitir_masivo_pdf', kwargs={'token': _token_del_correo()}
        ),
    )

    assert response.status_code == 200
    assert re.search(
        r'filename="76192083_dte_masivo_\d{14}_pdf.zip"',
        response['Content-Disposition'],
    )
    with zipfile.ZipFile(io.BytesIO(response.getvalue())) as zip_pdf:
        assert zip_pdf.namelist() == ['LibreDTE_76192083_33-0000007.pdf']


def test_un_enlace_alterado_no_descarga_nada(
    client: Client,
    contribuyente: Contribuyente,
    usuario: User,
) -> None:
    _procesar(contribuyente, usuario, _documento(), pdf=True)
    token = _token_del_correo()

    response = client.get(
        reverse('billing:emitir_masivo_pdf', kwargs={'token': token + 'x'}),
    )

    assert response.status_code == 404


def test_el_bloqueo_rechaza_la_misma_planilla_mientras_esta_tomado(
    contribuyente: Contribuyente,
) -> None:
    bulk_biller.acquire_lock(contribuyente.pk, 'huella')

    with pytest.raises(bulk_biller.BulkBillingInProgressError):
        bulk_biller.acquire_lock(contribuyente.pk, 'huella')
